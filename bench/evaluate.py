"""Official data-conversion check plus explicit OOXML checks for derived tasks."""

import hashlib, importlib.util, json, math, sys
from pathlib import Path
import openpyxl
from openpyxl.chart import LineChart

from common import runtime

BENCH = Path(runtime()["toolathlon"]) / "tasks/finalpool"


def source_rows():
    wb = openpyxl.load_workbook(
        BENCH / "excel-market-research/initial_workspace/Market_Data.xlsx",
        data_only=True,
    )
    return {
        r[0]: [r[1], r[3], r[4]]
        for r in wb["RawData"].values
        if isinstance(r[0], (int, float))
    }


def title_text(title):
    if title is None:
        return ""
    try:
        return "".join(run.t for p in title.tx.rich.p for run in p.r)
    except (AttributeError, TypeError):
        return ""


def normalized_ref(value):
    return str(value).replace("$", "").replace("'", "").lstrip("=")


def close(a, b):
    return (
        isinstance(a, (int, float))
        and not isinstance(a, bool)
        and math.isclose(a, b, rel_tol=1e-8, abs_tol=1e-6)
    )


def list_options(wb, formula):
    f = (formula or "").strip().lstrip("=")
    if f.startswith('"') and f.endswith('"'):
        return f[1:-1].split(",")
    if "!" not in f:
        item = wb.defined_names.get(f)
        f = item.attr_text if item else f
    if "!" in f:
        sheet, ref = f.split("!", 1)
        try:
            return [cell.value for row in wb[sheet.strip("'")][ref] for cell in row]
        except (KeyError, ValueError, TypeError):
            return []
    return []


def evaluate(task, workspace):
    workspace = Path(workspace)
    if task == "excel-market-research":
        output = workspace / "segment_growth_rates.xlsx"
        if output.is_file():
            a = openpyxl.load_workbook(output, read_only=True)
            g = openpyxl.load_workbook(
                BENCH / task / "groundtruth_workspace" / output.name, read_only=True
            )
            try:
                missing = [
                    c.value
                    for c in g.active[1]
                    if c.value not in [x.value for x in a.active[1]]
                ]
                if missing or a.active.max_row < g.active.max_row:
                    return {
                        "task": task,
                        "passed": False,
                        "detail": "Required columns or rows are missing",
                    }
            finally:
                a.close()
                g.close()
        path = BENCH / task / "evaluation/check_local.py"
        spec = importlib.util.spec_from_file_location("official_market", path)
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        passed, detail = mod.check_local(
            str(workspace), str(BENCH / task / "groundtruth_workspace")
        )
        return {"task": task, "passed": bool(passed), "detail": detail}
    if task == "excel-data-transformation":
        p = BENCH / task / "evaluation/check_local_improved.py"
        spec = importlib.util.spec_from_file_location("official_excel_conversion", p)
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        passed, detail = mod.check_local(
            str(workspace), str(BENCH / task / "groundtruth_workspace")
        )
        return {
            "task": task,
            "kind": "public-adapted",
            "passed": bool(passed),
            "detail": detail,
            "officialCheckSource": str(p),
        }
    filename = (
        "MarketTrends.xlsx" if task == "market-chart-report" else "MarketChecks.xlsx"
    )
    path = workspace / filename
    if not path.is_file():
        return {
            "task": task,
            "kind": "derived",
            "passed": False,
            "checks": {"file_exists": False},
        }
    wb = openpyxl.load_workbook(path, data_only=False)
    values = openpyxl.load_workbook(path, data_only=True)
    rows = source_rows()
    checks = {}
    original = BENCH / "excel-market-research/initial_workspace/Market_Data.xlsx"
    source = workspace / "Market_Data.xlsx"
    checks["source_unchanged"] = (
        source.is_file()
        and hashlib.sha256(source.read_bytes()).digest()
        == hashlib.sha256(original.read_bytes()).digest()
    )
    name = "Summary" if task == "market-chart-report" else "Audit"
    checks["sheet_exists"] = name in wb
    if name not in wb:
        return {"task": task, "kind": "derived", "passed": False, "checks": checks}
    s = wb[name]
    v = values[name]
    if task == "market-chart-report":
        checks["headers"] = [s.cell(1, c).value for c in range(1, 5)] == [
            "Year",
            "Agriculture",
            "Food",
            "Metals",
        ]
        checks["values"] = all(
            v.cell(i + 2, 1).value == year
            and all(close(v.cell(i + 2, j + 2).value, rows[year][j]) for j in range(3))
            for i, year in enumerate(range(2014, 2025))
        )
        # The task asks for one decimal place, not a specific grouping policy.
        # Accept the two established equivalent numeric formats; no percentage,
        # scaling, additional decimals, or arbitrary Excel-format inference.
        checks["number_format"] = all(
            s.cell(r, c).number_format in {"0.0", "#,##0.0"}
            for r in range(2, 13)
            for c in range(2, 5)
        )
        checks["header_bold"] = all(s.cell(1, c).font.bold for c in range(1, 5))
        checks["freeze_panes"] = s.freeze_panes == "B2"
        checks["print_layout"] = (
            s.page_setup.orientation == "landscape"
            and str(s.page_setup.paperSize) == "9"
            and s.page_setup.fitToWidth == 1
            and s.page_setup.fitToHeight == 0
            and bool(
                s.sheet_properties.pageSetUpPr
                and s.sheet_properties.pageSetUpPr.fitToPage
            )
        )
        checks["print_titles"] = normalized_ref(s.print_title_rows) == "1:1"
        checks["line_chart"] = len(s._charts) == 1 and isinstance(
            s._charts[0], LineChart
        )
        if checks["line_chart"]:
            chart = s._charts[0]
            checks["chart_title"] = title_text(chart.title) == "Market trends"
            checks["axis_titles"] = (
                title_text(chart.x_axis.title) == "Year"
                and title_text(chart.y_axis.title) == "Million USD"
            )
            checks["legend"] = chart.legend is not None and chart.legend.position == "b"
            checks["anchor"] = (
                chart.anchor._from.col == 5 and chart.anchor._from.row == 1
            )
            refs = []
            for ser in chart.ser:
                val = ser.val.numRef.f if ser.val and ser.val.numRef else ""
                cat = (
                    ser.cat.numRef.f
                    if ser.cat and ser.cat.numRef
                    else ser.cat.strRef.f
                    if ser.cat and ser.cat.strRef
                    else ""
                )
                refs.append((normalized_ref(val), normalized_ref(cat)))
            checks["series"] = sorted(refs) == [
                (f"Summary!{c}2:{c}12", "Summary!A2:A12") for c in "BCD"
            ]
    elif task == "market-validation-report":
        checks["headers"] = [s.cell(1, c).value for c in range(1, 6)] == [
            "Year",
            "AgricultureYoY",
            "FoodYoY",
            "MetalsYoY",
            "ReviewStatus",
        ]
        checks["values"] = all(
            v.cell(i + 2, 1).value == year
            and all(
                close(v.cell(i + 2, j + 2).value, rows[year][j] / rows[year - 1][j] - 1)
                for j in range(3)
            )
            for i, year in enumerate(range(2015, 2025))
        )
        checks["status_values"] = all(
            v.cell(r, 5).value == "Pending" for r in range(2, 12)
        )
        checks["number_format"] = all(
            s.cell(r, c).number_format == "0.0%"
            for r in range(2, 12)
            for c in range(2, 5)
        )
        checks["freeze_panes"] = s.freeze_panes == "A2"
        checks["validation"] = any(
            d.type == "list"
            and normalized_ref(d.sqref) == "E2:E11"
            and list_options(wb, d.formula1) == ["Pending", "Approved"]
            for d in s.data_validations.dataValidation
        )
        table = s.tables.get("GrowthAudit")
        checks["table"] = bool(
            table is not None
            and table.ref == "A1:E11"
            and table.tableStyleInfo
            and table.tableStyleInfo.name == "TableStyleMedium9"
            and table.tableStyleInfo.showRowStripes
        )
        checks["color_scale"] = False
        for cf, rules in s.conditional_formatting._cf_rules.items():
            if normalized_ref(cf.sqref) != "B2:D11":
                continue
            for rule in rules:
                if rule.type != "colorScale":
                    continue
                scale = rule.colorScale
                if len(scale.cfvo) != 3 or len(scale.color) != 3:
                    continue
                checks["color_scale"] |= (
                    [x.type for x in scale.cfvo] == ["min", "percentile", "max"]
                    and scale.cfvo[1].val == 50
                    and [str(c.rgb)[-6:].upper() for c in scale.color]
                    == ["FF0000", "FFFF00", "00FF00"]
                )
    else:
        raise ValueError("Unknown task")
    return {
        "task": task,
        "kind": "derived",
        "passed": all(checks.values()),
        "checks": checks,
        "boundary": "Programmatic workbook content and structure checks, not a rendered visual review. Derived task, not official benchmark score.",
    }


if __name__ == "__main__":
    try:
        result = evaluate(sys.argv[1], sys.argv[2])
    except Exception as error:
        result = {
            "task": sys.argv[1],
            "passed": False,
            "error": type(error).__name__ + ": " + str(error),
        }
    Path(sys.argv[3]).write_text(
        json.dumps(result, ensure_ascii=False, indent=2, default=str) + "\n"
    )
    print(json.dumps({"task": result["task"], "passed": result["passed"]}))
